import pandas as pd
import numpy as np
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

master_file = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\DRG versi September 2025-Final Publish & DRG Master.xlsx'
xl_m = pd.ExcelFile(master_file)
master_df = xl_m.parse('Rencana Usulan per 29092026', dtype=str)
dc_map_1 = master_df[['dc', 'dc desc']].dropna().rename(columns={'dc':'dc_code', 'dc desc':'dc_desc'})
dc_map_2 = master_df[['dc update', 'desc dc update']].dropna().rename(columns={'dc update':'dc_code', 'desc dc update':'dc_desc'})
dc_map = pd.concat([dc_map_1, dc_map_2]).drop_duplicates()
dc_map['dc_str'] = dc_map['dc_code'].str.replace(r'\.0$', '', regex=True).str.strip()
dc_map['dc'] = dc_map['dc_code']
dc_map['dc desc'] = dc_map['dc_desc']

def normalize_text(t):
    if pd.isna(t): return ''
    t = str(t).lower().strip()
    t = re.sub(r'[^a-z0-9]', '', t)
    return t

dc_map['desc_norm'] = dc_map['dc desc'].apply(normalize_text)

desc_to_dc = {}
for _, row in dc_map.iterrows():
    n = row['desc_norm']
    if n not in desc_to_dc:
        desc_to_dc[n] = {'dcs': set(), 'official_descs': set()}
    desc_to_dc[n]['dcs'].add(row['dc_str'])
    desc_to_dc[n]['official_descs'].add(row['dc desc'])

pdf_mapping = [
    ('P11AA', '11032, 11031', 'Endovascular Procedures for Neurologic Diseases'),
    ('P11AG', '11141', 'Complex Intracranial Procedures'),
    ('P11AH', '11071', 'Simple Intracranial Procedures'),
    ('P11AB', '11041', 'Complex Spinal Procedures'),
    ('P11AC', '11021', 'Complex Craniotomy Procedures'),
    ('P11AK', '11091', 'Simple Spinal Procedures'),
    ('P11AE', '11131', 'Moderate Craniotomy Procedures'),
    ('P11AF', '11051', 'Simple Craniotomy Procedures'),
    ('P11AJ', '11151', 'Moderate Spinal Procedures'),
    ('P11AP', '11161', 'Neuromodulation Procedures'),
    ('P11AN', '11101', 'Extracranial Vascular Procedures'),
    ('P11AL', '11081', 'Complex Peripheral Nerve Procedures'),
    ('P11AD', '11061', 'Ventricular System Procedures'),
    ('P11AM', '11121', 'Simple Peripheral Nerve Procedures'),
    ('P11AQ', '11171', 'Neurologic Angiography'),
    ('P11AR', '11182, 11181', 'Percutaneous Transluminal Angiography (PTA) Cerebral'),
    ('D11AB', '11621', 'Nervous System Neoplasm'),
    ('D11AP', '11741, 11742', 'Ischemic Stroke'),
    ('D11AN', '11731', 'Subarachnoid Hemorrhage'),
    ('D11AM', '11721', 'Hemorrhagic Stroke'),
    ('D11AD', '11642, 11641', 'Multiple Sclerosis & Other Nervous Diseases'),
    ('D11AL', '11711', 'Congenital Nervous System Disorders'),
    ('D11AA', '11611', 'Nervous System Infections'),
    ('D11AC', '11631', 'Systemic Atrophy'),
    ('D11AF', '11661', 'Neurobehaviour & Other Neurodegenerative Disease'),
    ('D11AU', '11791', 'Delirium & Coma'),
    ('D11AK', '11701', 'Cerebral Palsy & Others Paralytic Syndrome'),
    ('D11AH', '11681', 'Nerves, Radix, & Plexus Disorders'),
    ('D11AS', '11771', 'Medula Spinalis, Spine & Peripheral Nervous Trauma'),
    ('D11AV', '11801', 'Others Nervous System Disorders'),
    ('D11AJ', '11692', 'Myoneural Junction & Muscle Disease'),
    ('D11AQ', '11751', 'Transient Ischemic Attact & Others Neurovascular Disorders'),
    ('D11AR', '11761', 'Head Trauma'),
    ('D11AE', '11651', 'Extrapyramidal & Movement Disorders'),
    ('D11AG', '11671', 'Epilepsy & Convulsion Attack'),
    ('D11AT', '11781', 'Headache'),
    ('11CX', '11031', 'Endovascular Proc. for Complex Neurologic Diseases'),
    ('11PDX', '11742', 'Ischemic Stroke w Thrombolytic'),
    ('11PCX', '11641', 'Multiple Sclerosis & Other Nervous Diseases w Plasmapheresis Proc.'),
    ('11PEX', '11181', 'Percutaneous Transluminal Angiography (PTA) Cerebral w Stents Proc.'),
    ('11PFX', '11181', 'Percutaneous Transluminal Angiography (PTA) Cerebral w Stents Proc.')
]
for _, dc_val, desc_val in pdf_mapping:
    n = normalize_text(desc_val)
    if n not in desc_to_dc:
        desc_to_dc[n] = {'dcs': set(), 'official_descs': set()}
    for d in dc_val.split(','):
        desc_to_dc[n]['dcs'].add(d.strip())
    desc_to_dc[n]['official_descs'].add(desc_val)

def map_desc_to_dc(desc):
    if pd.isna(desc) or desc == '': return '', ''
    n = normalize_text(desc)
    if n in desc_to_dc:
        dcs = ', '.join(sorted(list(desc_to_dc[n]['dcs'])))
        off = ' | '.join(sorted(list(desc_to_dc[n]['official_descs'])))
        return dcs, off
    
    if 'intracranial' in n and 'vessel' in n:
        return '11031', 'Intracranial Vessel Proc. (Fallback)'
    return '', desc

file_path = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\MDC 11 Neuro - Cluster - Follow Up Usulan.xlsx'
xl = pd.ExcelFile(file_path)

def normalize_cols(df, code_col_variations):
    for col in code_col_variations:
        if col in df.columns:
            df = df.rename(columns={col: 'ICD_Code'})
            break
            
    df = df.dropna(subset=['ICD_Code']).copy()
    df['ICD_Code'] = df['ICD_Code'].astype(str).str.strip().str.upper()
    df['ICD_Code'] = df['ICD_Code'].replace({'NAN': np.nan, 'NONE': np.nan})
    df = df.dropna(subset=['ICD_Code']).copy()
    
    if 'New Cluster Description' in df.columns:
        df['Final_Desc'] = np.where(df['New Cluster Description'].notna() & (df['New Cluster Description'] != ''), df['New Cluster Description'], df['Cluster Description'])
    else:
        df['Final_Desc'] = df['Cluster Description']
        
    return df

def compare(sheet_web, sheet_usulan, icd_cols):
    df_web = xl.parse(sheet_web, dtype=str)
    df_usulan = xl.parse(sheet_usulan, dtype=str)
    
    df_web = normalize_cols(df_web, icd_cols)
    df_usulan = normalize_cols(df_usulan, icd_cols)
    
    web_sub = df_web[['ICD_Code', 'Final_Desc']].copy()
    web_sub.columns = ['ICD_Code', 'Desc_Web']
    
    usulan_sub = df_usulan[['ICD_Code', 'Final_Desc']].copy()
    usulan_sub.columns = ['ICD_Code', 'Desc_Usulan']
    
    merged = pd.merge(web_sub, usulan_sub, on='ICD_Code', how='outer')
    
    conditions = [
        merged['Desc_Web'].isna() | (merged['Desc_Web'] == ''),
        merged['Desc_Usulan'].isna() | (merged['Desc_Usulan'] == ''),
        merged.apply(lambda r: normalize_text(str(r['Desc_Web'])) != normalize_text(str(r['Desc_Usulan'])), axis=1)
    ]
    choices = ['Baru Ditambahkan', 'Dihapus (Drop)', 'Pindah DC']
    merged['Status'] = np.select(conditions, choices, default='Tetap')
    
    return merged

df_r10 = compare('ranap cluster icd10_web20260922', 'ranap clusters (ICD 10)', ['ICD-10 Code', 'ICD 10 Code'])
df_r9 = compare('ranap cluster icd9cm_web2026092', 'ranap clusters (ICD 9)', ['ICD-9-CM Code'])
df_j9 = compare('rajal cluster icd9cm_web2026092', 'rajal clusters (icd 9 cm)', ['ICD-9-CM Code'])

all_dfs = {'Ranap ICD 10': df_r10, 'Ranap ICD 9': df_r9, 'Rajal ICD 9': df_j9}
res = {}

for sheet, df in all_dfs.items():
    df[['DC_Awal', 'DC_Desc_Awal']] = df.apply(lambda row: pd.Series(map_desc_to_dc(row['Desc_Web'])), axis=1)
    df[['DC_Baru', 'DC_Desc_Baru']] = df.apply(lambda row: pd.Series(map_desc_to_dc(row['Desc_Usulan'])), axis=1)
    
    def get_warning(row):
        warn = []
        if pd.notna(row['Desc_Web']) and row['Desc_Web'] != '' and row['DC_Awal'] == '':
            warn.append(f"Deskripsi Awal tidak dikenal di DRG Master/Drawio: {row['Desc_Web']}")
        if pd.notna(row['Desc_Usulan']) and row['Desc_Usulan'] != '' and row['DC_Baru'] == '':
            warn.append(f"Deskripsi Baru tidak dikenal di DRG Master/Drawio: {row['Desc_Usulan']}")
        return ' | '.join(warn)
        
    df['Warning'] = df.apply(get_warning, axis=1)
    
    cols = ['ICD_Code', 'Status', 'DC_Awal', 'DC_Desc_Awal', 'DC_Baru', 'DC_Desc_Baru', 'Warning', 'Desc_Web', 'Desc_Usulan']
    df = df[cols]
    df = df[df['Status'] != 'Tetap']
    
    df = df.replace({'nan': np.nan, 'NAN': np.nan, 'None': np.nan})
    df = df.fillna('')
    res[sheet] = df

out_path = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\MDC_11_Analisis_Berdasarkan_Deskripsi.xlsx'
with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
    for sheet, df in res.items():
        df.to_excel(writer, sheet_name=sheet, index=False)
        
wb = openpyxl.load_workbook(out_path)
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    ws.freeze_panes = 'A2'
    fill_header = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    font_header = Font(color='FFFFFF', bold=True)
    for cell in ws[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.auto_filter.ref = ws.dimensions
    
    headers = {cell.value: cell.column for cell in ws[1]}
    status_col = headers.get('Status')
    warn_col = headers.get('Warning')
    
    for row in range(2, ws.max_row + 1):
        if status_col:
            c = ws.cell(row=row, column=status_col)
            if c.value == 'Baru Ditambahkan': c.fill, c.font = PatternFill(start_color='C6EFCE', fill_type='solid'), Font(color='006100')
            elif c.value == 'Dihapus (Drop)': c.fill, c.font = PatternFill(start_color='FFC7CE', fill_type='solid'), Font(color='9C0006')
            elif c.value == 'Pindah DC': c.fill, c.font = PatternFill(start_color='FFEB9C', fill_type='solid'), Font(color='9C6500')
            
        if warn_col:
            cw = ws.cell(row=row, column=warn_col)
            if cw.value and cw.value != '':
                cw.fill, cw.font = PatternFill(start_color='FF9999', fill_type='solid'), Font(color='000000', bold=True)

    for col in ws.columns:
        max_l = max([len(str(cell.value)) for cell in col if cell.value] + [0])
        ws.column_dimensions[col[0].column_letter].width = min(max_l + 2, 50)

wb.save(out_path)
print('Excel by Description Generated at:', out_path)
