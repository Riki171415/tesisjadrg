import pandas as pd
import numpy as np
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

# Master
master_file = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\DRG versi September 2025-Final Publish & DRG Master.xlsx'
xl_m = pd.ExcelFile(master_file)
master_df = xl_m.parse('Rencana Usulan per 29092026', dtype=str)

dc_map_1 = master_df[['dc', 'dc desc']].dropna().rename(columns={'dc':'dc_code', 'dc desc':'dc_desc'})
dc_map_2 = master_df[['dc update', 'desc dc update']].dropna().rename(columns={'dc update':'dc_code', 'desc dc update':'dc_desc'})
dc_map = pd.concat([dc_map_1, dc_map_2]).drop_duplicates()
dc_map['dc_str'] = dc_map['dc_code'].str.replace(r'\.0$', '', regex=True).str.strip()
dc_dict = dict(zip(dc_map['dc_str'], dc_map['dc_desc']))

# MDC 11 mapping
pdf_mapping = [
    ('P11AA', '11032, 11031', 'Endovascular Procedures for Neurologic Diseases'),
    ('P11AG', '11141', 'Complex Intracranial Procedures'),
    ('P11AH', '11071', 'Simple Intracranial Procedures'),
    ('P11AB', '11041', 'Complex Spinal Procedures'),
    ('P11AC', '11021', 'Complex Craniotomy Procedures'),
    ('P11AK', '11091', 'Simple Spinal Procedures'),
    ('P11AE', '11131', 'Moderate Craniotomy Procedures'),
    ('P11AF', '11051', 'Simple Craniotomy Procedures'),
    ('P11AJ', '11151', 'Moderate Spinal Procedures'),
    ('P11AP', '11161', 'Neuromodulation Procedures'),
    ('P11AN', '11101', 'Extracranial Vascular Procedures'),
    ('P11AL', '11081', 'Complex Peripheral Nerve Procedures'),
    ('P11AD', '11061', 'Ventricular System Procedures'),
    ('P11AM', '11121', 'Simple Peripheral Nerve Procedures'),
    ('P11AQ', '11171', 'Neurologic Angiography'),
    ('P11AR', '11182, 11181', 'Percutaneous Transluminal Angiography (PTA) Cerebral'),
    ('D11AB', '11621', 'Nervous System Neoplasm'),
    ('D11AP', '11741, 11742', 'Ischemic Stroke'),
    ('D11AN', '11731', 'Subarachnoid Hemorrhage'),
    ('D11AM', '11721', 'Hemorrhagic Stroke'),
    ('D11AD', '11642, 11641', 'Multiple Sclerosis & Other Nervous Diseases'),
    ('D11AL', '11711', 'Congenital Nervous System Disorders'),
    ('D11AA', '11611', 'Nervous System Infections'),
    ('D11AC', '11631', 'Systemic Atrophy'),
    ('D11AF', '11661', 'Neurobehaviour & Other Neurodegenerative Disease'),
    ('D11AU', '11791', 'Delirium & Coma'),
    ('D11AK', '11701', 'Cerebral Palsy & Others Paralytic Syndrome'),
    ('D11AH', '11681', 'Nerves, Radix, & Plexus Disorders'),
    ('D11AS', '11771', 'Medula Spinalis, Spine & Peripheral Nervous Trauma'),
    ('D11AV', '11801', 'Others Nervous System Disorders'),
    ('D11AJ', '11692', 'Myoneural Junction & Muscle Disease'),
    ('D11AQ', '11751', 'Transient Ischemic Attact & Others Neurovascular Disorders'),
    ('D11AR', '11761', 'Head Trauma'),
    ('D11AE', '11651', 'Extrapyramidal & Movement Disorders'),
    ('D11AG', '11671', 'Epilepsy & Convulsion Attack'),
    ('D11AT', '11781', 'Headache'),
]
drawio_dc_dict = {p[0]: p[1] for p in pdf_mapping}
drawio_desc_dict = {p[0]: p[2] for p in pdf_mapping}

file_path = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\MDC 11 Neuro - Cluster - Follow Up Usulan.xlsx'
xl = pd.ExcelFile(file_path)

def normalize_cols(df, code_col_variations):
    for col in code_col_variations:
        if col in df.columns:
            df = df.rename(columns={col: 'ICD_Code'})
            break
            
    if 'New Cluster Code' in df.columns:
        df['Final_Cluster_Code'] = np.where(df['New Cluster Code'].notna(), df['New Cluster Code'], df['Cluster Code'])
    else:
        df['Final_Cluster_Code'] = df['Cluster Code']
        
    if 'New Cluster Description' in df.columns:
        df['Final_Desc'] = np.where(df['New Cluster Description'].notna() & (df['New Cluster Description'] != ''), df['New Cluster Description'], df['Cluster Description'])
    else:
        df['Final_Desc'] = df['Cluster Description']
    
    df['ICD_Code'] = df['ICD_Code'].astype(str).str.strip().str.upper()
    df['Final_Cluster_Code'] = df['Final_Cluster_Code'].astype(str).str.strip().str.upper()
    df['ICD_Code'] = df['ICD_Code'].replace({'NAN': np.nan, 'NONE': np.nan})
    df['Final_Cluster_Code'] = df['Final_Cluster_Code'].replace({'NAN': np.nan, 'NONE': np.nan})
    df = df.dropna(subset=['ICD_Code']).copy()
    
    return df

def is_consistent(excel_desc, drawio_desc):
    if pd.isna(excel_desc) or pd.isna(drawio_desc) or str(excel_desc) == 'nan' or str(drawio_desc) == 'nan': return ''
    e = re.sub(r'[^a-z0-9]', '', str(excel_desc).lower().strip())
    d = re.sub(r'[^a-z0-9]', '', str(drawio_desc).lower().strip())
    if e == d: return ''
    return 'TIDAK SINKRON: Excel ("' + str(excel_desc) + '") vs Draw.io ("' + str(drawio_desc) + '")'

def compare(sheet_web, sheet_usulan, icd_cols):
    df_web = xl.parse(sheet_web, dtype=str)
    df_usulan = xl.parse(sheet_usulan, dtype=str)
    
    df_web = normalize_cols(df_web, icd_cols)
    df_usulan = normalize_cols(df_usulan, icd_cols)
    
    web_sub = df_web[['ICD_Code', 'Final_Cluster_Code', 'Final_Desc']].copy()
    web_sub.columns = ['ICD_Code', 'PDC_Awal', 'Desc_Awal']
    
    usulan_sub = df_usulan[['ICD_Code', 'Final_Cluster_Code', 'Final_Desc']].copy()
    usulan_sub.columns = ['ICD_Code', 'PDC_Baru', 'Desc_Baru']
    
    merged = pd.merge(web_sub, usulan_sub, on='ICD_Code', how='outer')
    
    conditions = [
        merged['PDC_Awal'].isna(),
        merged['PDC_Baru'].isna(),
        merged['PDC_Awal'] != merged['PDC_Baru']
    ]
    choices = ['Baru Ditambahkan', 'Dihapus (Drop)', 'Pindah DC/Cluster']
    merged['Status_Perubahan'] = np.select(conditions, choices, default='Tetap')
    
    merged['DC_Awal'] = merged['PDC_Awal'].map(drawio_dc_dict)
    merged['DC_Baru'] = merged['PDC_Baru'].map(drawio_dc_dict)
    
    merged['Drawio_Desc_Awal'] = merged['PDC_Awal'].map(drawio_desc_dict)
    merged['Drawio_Desc_Baru'] = merged['PDC_Baru'].map(drawio_desc_dict)
    
    merged['Sync_Awal'] = merged.apply(lambda r: is_consistent(r['Desc_Awal'], r['Drawio_Desc_Awal']) if pd.notna(r['PDC_Awal']) else '', axis=1)
    merged['Sync_Baru'] = merged.apply(lambda r: is_consistent(r['Desc_Baru'], r['Drawio_Desc_Baru']) if pd.notna(r['PDC_Baru']) else '', axis=1)
    
    merged['Ketidaksinkronan_Data'] = merged[['Sync_Awal', 'Sync_Baru']].apply(lambda x: ' | '.join([i for i in x if i and i != '']), axis=1)
    
    return merged[merged['Status_Perubahan'] != 'Tetap'].copy()

df_r10 = compare('ranap cluster icd10_web20260922', 'ranap clusters (ICD 10)', ['ICD-10 Code', 'ICD 10 Code'])
df_r9 = compare('ranap cluster icd9cm_web2026092', 'ranap clusters (ICD 9)', ['ICD-9-CM Code'])
df_j9 = compare('rajal cluster icd9cm_web2026092', 'rajal clusters (icd 9 cm)', ['ICD-9-CM Code'])

all_changes = [df_r10, df_r9, df_j9]

dc_stats = {}
for df in all_changes:
    for _, row in df.iterrows():
        dc_awal_list = [d.strip() for d in str(row['DC_Awal']).split(',')] if pd.notna(row['DC_Awal']) and str(row['DC_Awal']) != 'nan' else []
        dc_baru_list = [d.strip() for d in str(row['DC_Baru']).split(',')] if pd.notna(row['DC_Baru']) and str(row['DC_Baru']) != 'nan' else []
        
        stat = row['Status_Perubahan']
        if stat == 'Baru Ditambahkan':
            for d in dc_baru_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Masuk'] += 1
        elif stat == 'Dihapus (Drop)':
            for d in dc_awal_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Keluar'] += 1
        elif stat == 'Pindah DC/Cluster':
            for d in dc_awal_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Keluar'] += 1
            for d in dc_baru_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Masuk'] += 1

summary_data = []
for dc, counts in dc_stats.items():
    desc = dc_dict.get(dc, 'Tidak Ditemukan di DRG Master')
    summary_data.append({
        'Nomor DC Terkait (Draw.io)': dc,
        'Nama DC (DRG Master)': desc,
        'Jumlah ICD Baru / Pindah Masuk (+)': counts['Masuk'],
        'Jumlah ICD Drop / Pindah Keluar (-)': counts['Keluar']
    })
df_summary = pd.DataFrame(summary_data).sort_values('Nomor DC Terkait (Draw.io)')

out_path = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\MDC_11_Terstruktur_DC_V2.xlsx'
with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
    df_summary.to_excel(writer, sheet_name='1. Ringkasan DC Terdampak', index=False)
    
    sheets = [('2. Detail Ranap ICD 10', df_r10), ('3. Detail Ranap ICD 9', df_r9), ('4. Detail Rajal ICD 9', df_j9)]
    cols = ['ICD_Code', 'Status_Perubahan', 'DC_Awal', 'PDC_Awal', 'Desc_Awal', 'DC_Baru', 'PDC_Baru', 'Desc_Baru', 'Ketidaksinkronan_Data']
    
    for name, df in sheets:
        df_out = df[cols].copy()
        df_out = df_out.replace({'nan': np.nan, 'NAN': np.nan, 'None': np.nan})
        df_out = df_out.fillna('')
        df_out.to_excel(writer, sheet_name=name, index=False)

wb = openpyxl.load_workbook(out_path)
fill_err = PatternFill(start_color='FFC7CE', fill_type='solid')
font_err = Font(color='9C0006', bold=True)
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill(start_color='4F81BD', fill_type='solid')
        cell.font = Font(color='FFFFFF', bold=True)
    headers = {cell.value: cell.column for cell in ws[1]}
    stat_col = headers.get('Status_Perubahan')
    warn_col = headers.get('Ketidaksinkronan_Data')
    for row in range(2, ws.max_row + 1):
        if stat_col:
            c = ws.cell(row=row, column=stat_col)
            if c.value and 'Baru' in str(c.value): c.fill = PatternFill(start_color='C6EFCE', fill_type='solid')
            elif c.value and 'Dihapus' in str(c.value): c.fill = PatternFill(start_color='FFC7CE', fill_type='solid')
            elif c.value and 'Pindah' in str(c.value): c.fill = PatternFill(start_color='FFEB9C', fill_type='solid')
        if warn_col:
            c = ws.cell(row=row, column=warn_col)
            if c.value and c.value != '': c.fill, c.font = fill_err, font_err
    for col in ws.columns:
        max_l = max([len(str(cell.value)) for cell in col if cell.value] + [0])
        ws.column_dimensions[col[0].column_letter].width = min(max_l + 2, 50)
wb.save(out_path)
print('Done!')
