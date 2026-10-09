import pandas as pd
import numpy as np
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

# Master
master_file = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\DRG versi September 2025-Final Publish & DRG Master.xlsx'
xl_m = pd.ExcelFile(master_file)
master_df = xl_m.parse('Rencana Usulan per 29092026', dtype=str)

dc_map_1 = master_df[['dc', 'dc desc']].dropna().rename(columns={'dc':'dc_code', 'dc desc':'dc_desc'})
dc_map_2 = master_df[['dc update', 'desc dc update']].dropna().rename(columns={'dc update':'dc_code', 'desc dc update':'dc_desc'})
dc_map = pd.concat([dc_map_1, dc_map_2]).drop_duplicates()
dc_map['dc_str'] = dc_map['dc_code'].str.replace(r'\.0$', '', regex=True).str.strip()
dc_dict = dict(zip(dc_map['dc_str'], dc_map['dc_desc']))

def get_dc_desc(dcs):
    if pd.isna(dcs) or str(dcs) == 'nan' or not dcs: return ''
    lst = [d.strip() for d in str(dcs).split(',')]
    res = [dc_dict.get(d, 'Tidak Terdaftar di Master') for d in lst]
    return ' | '.join(res)

# MDC 12 mapping
pdf_mapping = [
    ('P12AA', '12011', 'Retinal Procedures'),
    ('P12AB', '12021', 'Major Corneal, Scleral & Conjunctival Procedures'),
    ('P12AC', '12031', 'Glaucoma Procedures'),
    ('P12AD', '12041', 'Lens Procedures'),
    ('P12AE', '12051', 'Eyelid and Lacrymal Procedures'),
    ('P12AG', '12071', 'Enucleations & Orbital Procedures'),
    ('P12AH', '12081', 'Strabismus Procedures'),
    ('P12AJ', '12091', 'Minor Corneal, Scleral & Conjunctival Procedures'),
    ('P12AK', '12101', 'Others Eye Procedures'),
    ('P12AL', '12411', 'Non Surgical Eye Procedures'),
    ('D12AA', '12691', 'Acute Infections or Inflammation'),
    ('D12AC', '12611', 'Hyphema & Medically Managed Trauma'),
    ('D12AD', '12621', 'Glaucoma Disorders'),
    ('D12AE', '12631', 'Choroid, Vitreus & Retinal Disorders'),
    ('D12AF', '12641', 'Lens & Corneal Disorder'),
    ('D12AH', '12651', 'Eyelid & Conjunctival Disorder'),
    ('D12AK', '12661', 'Other Disorder of the Eye'),
    ('D12AL', '12681, 12701', 'Eye Neoplasms'),
    ('D12AJ', '12681, 12701', 'Eye Neoplasms'),
]
drawio_dc_dict = {p[0]: p[1] for p in pdf_mapping}
drawio_desc_dict = {p[0]: p[2] for p in pdf_mapping}

file_path = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\MDC 12 Eye - Cluster 20260901 - Follow Up Usulan.xlsx'
xl = pd.ExcelFile(file_path)

def normalize_cols(df, code_col_variations):
    for col in code_col_variations:
        if col in df.columns:
            df = df.rename(columns={col: 'ICD_Code'})
            break
            
    if 'New Cluster Code' in df.columns:
        df['Final_Cluster_Code'] = np.where(df['New Cluster Code'].notna(), df['New Cluster Code'], df['Cluster Code'])
    else:
        df['Final_Cluster_Code'] = df['Cluster Code']
        
    if 'New Cluster Description' in df.columns:
        df['Final_Desc'] = np.where(df['New Cluster Description'].notna() & (df['New Cluster Description'] != ''), df['New Cluster Description'], df['Cluster Description'])
    else:
        df['Final_Desc'] = df['Cluster Description']
        
    # NEW RULE: Check for 'Action' == 'drop'
    if 'Action' in df.columns:
        df['Final_Cluster_Code'] = np.where(df['Action'].astype(str).str.lower().str.strip() == 'drop', np.nan, df['Final_Cluster_Code'])
    
    df['ICD_Code'] = df['ICD_Code'].astype(str).str.strip().str.upper()
    df['Final_Cluster_Code'] = df['Final_Cluster_Code'].astype(str).str.strip().str.upper()
    df['ICD_Code'] = df['ICD_Code'].replace({'NAN': np.nan, 'NONE': np.nan})
    df['Final_Cluster_Code'] = df['Final_Cluster_Code'].replace({'NAN': np.nan, 'NONE': np.nan})
    df = df.dropna(subset=['ICD_Code']).copy()
    return df

def is_consistent(excel_desc, drawio_desc):
    if pd.isna(excel_desc) or pd.isna(drawio_desc) or str(excel_desc) == 'nan' or str(drawio_desc) == 'nan': return ''
    e = re.sub(r'[^a-z0-9]', '', str(excel_desc).lower().strip())
    d = re.sub(r'[^a-z0-9]', '', str(drawio_desc).lower().strip())
    if e == d: return ''
    return 'TIDAK SINKRON: Excel ("' + str(excel_desc) + '") vs Draw.io ("' + str(drawio_desc) + '")'

def compare(sheet_web, sheet_usulan, icd_cols):
    df_web = xl.parse(sheet_web, dtype=str)
    df_usulan = xl.parse(sheet_usulan, dtype=str)
    
    df_web = normalize_cols(df_web, icd_cols)
    df_usulan = normalize_cols(df_usulan, icd_cols)
    
    web_sub = df_web[['ICD_Code', 'Final_Cluster_Code', 'Final_Desc']].copy()
    web_sub.columns = ['ICD_Code', 'PDC_Awal', 'Desc_Awal_Excel']
    
    usulan_sub = df_usulan[['ICD_Code', 'Final_Cluster_Code', 'Final_Desc']].copy()
    usulan_sub.columns = ['ICD_Code', 'PDC_Baru', 'Desc_Baru_Excel']
    
    merged = pd.merge(web_sub, usulan_sub, on='ICD_Code', how='outer')
    
    conditions = [
        merged['PDC_Awal'].isna(),
        merged['PDC_Baru'].isna(),
        merged['PDC_Awal'] != merged['PDC_Baru']
    ]
    choices = ['Baru Ditambahkan', 'Dihapus (Drop)', 'Pindah DC/Cluster']
    merged['Status_Perubahan'] = np.select(conditions, choices, default='Tetap')
    
    merged['DC_Awal'] = merged['PDC_Awal'].map(drawio_dc_dict)
    merged['Nama_DC_Awal'] = merged['DC_Awal'].apply(get_dc_desc)
    
    merged['DC_Baru'] = merged['PDC_Baru'].map(drawio_dc_dict)
    merged['Nama_DC_Baru'] = merged['DC_Baru'].apply(get_dc_desc)
    
    merged['Drawio_Desc_Awal'] = merged['PDC_Awal'].map(drawio_desc_dict)
    merged['Drawio_Desc_Baru'] = merged['PDC_Baru'].map(drawio_desc_dict)
    
    merged['Sync_Awal'] = merged.apply(lambda r: is_consistent(r['Desc_Awal_Excel'], r['Drawio_Desc_Awal']) if pd.notna(r['PDC_Awal']) else '', axis=1)
    merged['Sync_Baru'] = merged.apply(lambda r: is_consistent(r['Desc_Baru_Excel'], r['Drawio_Desc_Baru']) if pd.notna(r['PDC_Baru']) else '', axis=1)
    merged['Ketidaksinkronan_Data'] = merged[['Sync_Awal', 'Sync_Baru']].apply(lambda x: ' | '.join([i for i in x if i and i != '']), axis=1)
    
    return merged[merged['Status_Perubahan'] != 'Tetap'].copy()

df_r10 = compare('ranap cluster icd10_web20260922', 'ranap clusters (ICD 10)', ['ICD-10 Code', 'ICD 10 Code'])
df_r9 = compare('ranap cluster icd9cm_web2026092', 'ranap clusters (ICD 9)', ['ICD-9-CM Code'])
df_j9 = compare('rajal cluster icd9cm_web2026092', 'rajal clusters (ICD 9)', ['ICD-9-CM Code'])

all_changes = [df_r10, df_r9, df_j9]
dc_stats = {}
for df in all_changes:
    for _, row in df.iterrows():
        dc_awal_list = [d.strip() for d in str(row['DC_Awal']).split(',')] if pd.notna(row['DC_Awal']) and str(row['DC_Awal']) != 'nan' else []
        dc_baru_list = [d.strip() for d in str(row['DC_Baru']).split(',')] if pd.notna(row['DC_Baru']) and str(row['DC_Baru']) != 'nan' else []
        stat = row['Status_Perubahan']
        if stat == 'Baru Ditambahkan':
            for d in dc_baru_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Masuk'] += 1
        elif stat == 'Dihapus (Drop)':
            for d in dc_awal_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Keluar'] += 1
        elif stat == 'Pindah DC/Cluster':
            for d in dc_awal_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Keluar'] += 1
            for d in dc_baru_list:
                if not d: continue
                if d not in dc_stats: dc_stats[d] = {'Masuk':0, 'Keluar':0}
                dc_stats[d]['Masuk'] += 1

summary_data = []
for dc, counts in dc_stats.items():
    desc = dc_dict.get(dc, 'Tidak Ditemukan di DRG Master')
    summary_data.append({
        'Nomor DC Terkait (Draw.io)': dc,
        'Nama DC (DRG Master)': desc,
        'Jumlah ICD Baru / Pindah Masuk (+)': counts['Masuk'],
        'Jumlah ICD Drop / Pindah Keluar (-)': counts['Keluar']
    })
df_summary = pd.DataFrame(summary_data).sort_values('Nomor DC Terkait (Draw.io)')

out_path = r'C:\Users\PUSBIKES-KEMKES\Downloads\mdc\MDC_12_Terstruktur_DC_V2.xlsx'
with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
    df_summary.to_excel(writer, sheet_name='1. Ringkasan DC Terdampak', index=False)
    
    sheets = [('2. Detail Ranap ICD 10', df_r10), ('3. Detail Ranap ICD 9', df_r9), ('4. Detail Rajal ICD 9', df_j9)]
    cols = ['ICD_Code', 'Status_Perubahan', 
            'DC_Awal', 'Nama_DC_Awal', 'PDC_Awal', 'Desc_Awal_Excel', 
            'DC_Baru', 'Nama_DC_Baru', 'PDC_Baru', 'Desc_Baru_Excel', 
            'Ketidaksinkronan_Data']
    
    for name, df in sheets:
        df_out = df[cols].copy()
        df_out = df_out.replace({'nan': np.nan, 'NAN': np.nan, 'None': np.nan})
        df_out = df_out.fillna('')
        df_out.to_excel(writer, sheet_name=name, index=False)

wb = openpyxl.load_workbook(out_path)
fill_err = PatternFill(start_color='FFC7CE', fill_type='solid')
font_err = Font(color='9C0006', bold=True)
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill(start_color='4F81BD', fill_type='solid')
        cell.font = Font(color='FFFFFF', bold=True)
    headers = {cell.value: cell.column for cell in ws[1]}
    stat_col = headers.get('Status_Perubahan')
    warn_col = headers.get('Ketidaksinkronan_Data')
    for row in range(2, ws.max_row + 1):
        if stat_col:
            c = ws.cell(row=row, column=stat_col)
            if c.value and 'Baru' in str(c.value): c.fill = PatternFill(start_color='C6EFCE', fill_type='solid')
            elif c.value and 'Dihapus' in str(c.value): c.fill = PatternFill(start_color='FFC7CE', fill_type='solid')
            elif c.value and 'Pindah' in str(c.value): c.fill = PatternFill(start_color='FFEB9C', fill_type='solid')
        if warn_col:
            c = ws.cell(row=row, column=warn_col)
            if c.value and c.value != '': c.fill, c.font = fill_err, font_err
    for col in ws.columns:
        max_l = max([len(str(cell.value)) for cell in col if cell.value] + [0])
        ws.column_dimensions[col[0].column_letter].width = min(max_l + 2, 50)
wb.save(out_path)
print('Done!')
